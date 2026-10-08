from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from crm.models import FinancialAccount, FinancialTransaction

SHEET_NAME = "Statement of Account"
REQUIRED_COLUMNS = (
    "Date",
    "Account N",
    "Debit",
    "Credit",
    "Entry Comment",
    "Operation ID",
    "Sender Name",
    "Sender Account N",
    "Recipient Name",
    "Recipient Account N",
    "Nomination",
    "Amount",
)


@dataclass
class ParsedBogRow:
    account_iban: str
    date: date
    amount: Decimal
    kind: str
    income_type: str
    counterparty_name: str
    counterparty_iban: str
    description: str
    external_id: str


@dataclass
class BogStatementImportResult:
    created: int
    skipped: int
    error: str


def import_bog_statement(file) -> BogStatementImportResult:
    rows, error = parse_bog_statement(file)
    if error:
        return BogStatementImportResult(created=0, skipped=0, error=error)

    ibans = {row.account_iban for row in rows}
    found = list(FinancialAccount.objects.filter(iban__in=ibans))
    accounts: dict[str, FinancialAccount] = {}
    ambiguous: list[str] = []
    for account in found:
        if account.iban in accounts:
            ambiguous.append(account.iban)
        else:
            accounts[account.iban] = account
    missing = sorted(ibans - set(accounts))
    ambiguous = sorted(set(ambiguous))
    messages: list[str] = []
    if missing:
        messages.append(_("Unknown accounts: %(accounts)s") % {"accounts": ", ".join(missing)})
    if ambiguous:
        messages.append(_("Accounts matched more than once: %(accounts)s") % {"accounts": ", ".join(ambiguous)})
    if messages:
        return BogStatementImportResult(created=0, skipped=0, error="; ".join(messages))

    external_ids = [row.external_id for row in rows]
    account_ids = [account.id for account in accounts.values()]
    with transaction.atomic():
        existing = set(
            FinancialTransaction.objects.filter(
                account_id__in=account_ids,
                external_id__in=external_ids,
            ).values_list("account_id", "external_id")
        )
        to_create: list[FinancialTransaction] = []
        skipped = 0
        for row in rows:
            account = accounts[row.account_iban]
            key = (account.id, row.external_id)
            if key in existing:
                skipped += 1
                continue
            existing.add(key)
            to_create.append(
                FinancialTransaction(
                    account=account,
                    date=row.date,
                    amount=row.amount,
                    kind=row.kind,
                    income_type=row.income_type,
                    counterparty_name=row.counterparty_name,
                    counterparty_iban=row.counterparty_iban,
                    description=row.description,
                    external_id=row.external_id,
                )
            )
        FinancialTransaction.objects.bulk_create(to_create)
    return BogStatementImportResult(created=len(to_create), skipped=skipped, error="")


def parse_bog_statement(file) -> tuple[list[ParsedBogRow], str]:
    workbook = load_workbook(file, data_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        return [], _("Sheet 'Statement of Account' not found")
    worksheet = workbook[SHEET_NAME]
    header_map: dict[str, int] | None = None
    rows: list[ParsedBogRow] = []
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        if header_map is None:
            names: dict[str, int] = {}
            for index, value in enumerate(values):
                if isinstance(value, str) and value.strip():
                    names[value.strip()] = index
            if "Date" in names and "Account N" in names:
                missing_columns = [name for name in REQUIRED_COLUMNS if name not in names]
                if missing_columns:
                    return [], _("Missing columns: %(columns)s") % {"columns": ", ".join(missing_columns)}
                header_map = names
            continue
        if not any(is_filled(value) for value in values):
            continue
        row_number = excel_row[0].row
        parsed, error = parse_bog_row(values, header_map, row_number)
        if error:
            return [], error
        rows.append(parsed)
    if header_map is None:
        return [], _("Statement header not found")
    return rows, ""


def parse_bog_row(
    values: list,
    header_map: dict[str, int],
    row_number: int,
) -> tuple[ParsedBogRow | None, str]:
    external_id = text_value(column_value(values, header_map, "Operation ID"))
    if external_id == "":
        return None, _("Row %(row_number)s: missing Operation ID") % {"row_number": row_number}
    has_debit = is_filled(column_value(values, header_map, "Debit"))
    has_credit = is_filled(column_value(values, header_map, "Credit"))
    if has_debit and has_credit:
        return None, _("Row %(row_number)s: both debit and credit are set") % {"row_number": row_number}
    if not has_debit and not has_credit:
        return None, _("Row %(row_number)s: debit and credit are empty") % {"row_number": row_number}
    raw_amount = column_value(values, header_map, "Amount")
    if not is_filled(raw_amount):
        return None, _("Row %(row_number)s: missing Amount") % {"row_number": row_number}
    amount = decimal_amount(raw_amount)
    if has_credit and amount <= 0:
        return None, _("Row %(row_number)s: amount sign does not match credit") % {"row_number": row_number}
    if has_debit and amount >= 0:
        return None, _("Row %(row_number)s: amount sign does not match debit") % {"row_number": row_number}
    parsed_date = parse_date(column_value(values, header_map, "Date"))
    if parsed_date is None:
        return None, _("Row %(row_number)s: invalid date") % {"row_number": row_number}
    account_iban = text_value(column_value(values, header_map, "Account N"))
    if account_iban == "":
        return None, _("Row %(row_number)s: missing Account N") % {"row_number": row_number}
    nomination = text_value(column_value(values, header_map, "Nomination"))
    if nomination:
        description = nomination
    else:
        description = text_value(column_value(values, header_map, "Entry Comment"))
    if has_credit:
        kind = FinancialTransaction.KIND_INCOME
        counterparty_name = text_value(column_value(values, header_map, "Sender Name"))
        counterparty_iban = text_value(column_value(values, header_map, "Sender Account N"))
    else:
        kind = FinancialTransaction.KIND_EXPENSE
        counterparty_name = text_value(column_value(values, header_map, "Recipient Name"))
        counterparty_iban = text_value(column_value(values, header_map, "Recipient Account N"))
    return (
        ParsedBogRow(
            account_iban=account_iban,
            date=parsed_date,
            amount=amount,
            kind=kind,
            income_type=FinancialTransaction.INCOME_TRANSFER,
            counterparty_name=counterparty_name,
            counterparty_iban=counterparty_iban,
            description=description,
            external_id=external_id,
        ),
        "",
    )


def column_value(values: list, header_map: dict[str, int], name: str):
    index = header_map[name]
    if index >= len(values):
        return None
    return values[index]


def is_filled(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str) and value.strip() == "":
        return False
    return True


def text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def decimal_amount(value) -> Decimal:
    if isinstance(value, Decimal):
        amount = value
    else:
        amount = Decimal(str(value))
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=int(value))).date()
    return None
