from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook

from crm.bog_personal_statement import import_bog_personal_statement
from crm.models import (
    FinancialAccount,
    FinancialTransaction,
    FinancialTransactionRule,
    FinancialTransactionRuleCondition,
)

ACCOUNT_NO = "GE25BG0000000609431766"
GEL_IBAN = f"{ACCOUNT_NO}GEL"
USD_IBAN = f"{ACCOUNT_NO}USD"
HEADERS = ["Date", "Details", None, "GEL", "USD", "EUR", "GBP"]
MERCHANT_DETAILS = (
    " Payment - Amount: GEL5.00; Merchant: Batumi Airport, Batumi, Airport Highway; "
    "MCC:5511; Date: 02/09/2026 22:41; Card No: ****3319; "
    "Payment transaction amount and currency: 5.00 GEL"
)
INCOME_CONVERSION_DETAILS = " Income - Amount GEL51.70; Automatic conversion, rate: 2.563"
EXPENSE_CONVERSION_DETAILS = " Payment - Amount USD20.17; Automatic conversion, rate: 2.563"
OUTGOING_DETAILS = (
    " Outgoing Transfer - Amount: GEL200.00; Beneficiary: im dzmitri haravets; "
    "Account: GE13BG0000000609702806GEL; Bank: JSC Bank of Georgia BAGAGE22; Details: Transfer"
)
INCOMING_DETAILS = (
    " Incoming Transfer - Amount: GEL1,350.00; Sender: i/m daria kharlamova; "
    "Account: GE07TB7154245068100036; Bank: JSC TBC Bank TBCBGE22; Details: Bonus"
)
WITHDRAWAL_DETAILS = (
    " Withdrawal - Amount: GEL1,000.00; ATM: Bank of Georgia, Batumi, Kobaladze St. 15/73 (2); "
    "MCC:6011; Date: 11/09/2026 11:18; Card No: ****3319; "
    "Payment transaction amount and currency: 1,000.00 GEL"
)


def _statement(rows: list[dict]) -> BytesIO:
    workbook = Workbook()
    details = workbook.active
    details.title = "Details"
    details["A3"] = "Account No:"
    details["C3"] = ACCOUNT_NO
    worksheet = workbook.create_sheet("Transactions")
    for index, header in enumerate(HEADERS, start=1):
        if header is not None:
            worksheet.cell(1, index, header)
    for offset, row in enumerate(rows):
        excel_row = 2 + offset
        worksheet.cell(excel_row, 1, row["Date"])
        worksheet.cell(excel_row, 2, row["Details"])
        if "GEL" in row:
            worksheet.cell(excel_row, 4, row["GEL"])
        if "USD" in row:
            worksheet.cell(excel_row, 5, row["USD"])
        if "EUR" in row:
            worksheet.cell(excel_row, 6, row["EUR"])
        if "GBP" in row:
            worksheet.cell(excel_row, 7, row["GBP"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _account(iban: str, currency: str) -> FinancialAccount:
    return FinancialAccount.objects.create(
        name=f"BOG personal {currency}",
        kind=FinancialAccount.KIND_BANK,
        iban=iban,
        currency=currency,
    )


def _rows() -> list[dict]:
    return [
        {"Date": "02/09/2026", "Details": MERCHANT_DETAILS, "GEL": -5},
        {"Date": "02/09/2026", "Details": INCOME_CONVERSION_DETAILS, "GEL": 51.7},
        {"Date": "02/09/2026", "Details": EXPENSE_CONVERSION_DETAILS, "USD": -20.17},
        {"Date": "06/09/2026", "Details": OUTGOING_DETAILS, "GEL": -200},
        {"Date": "08/09/2026", "Details": INCOMING_DETAILS, "GEL": 1350},
        {"Date": "11/09/2026", "Details": WITHDRAWAL_DETAILS, "GEL": -1000},
    ]


def _external_id(day: date, currency: str, amount: Decimal, details: str) -> str:
    return f"{day.isoformat()}|{currency}|{amount}|{details.strip()}"


class BogPersonalStatementImportTests(TestCase):
    def test_creates_card_payment_conversion_transfers_and_withdrawal(self):
        gel = _account(GEL_IBAN, "GEL")
        usd = _account(USD_IBAN, "USD")
        result = import_bog_personal_statement(_statement(_rows()))
        self.assertEqual(result.error, "")
        self.assertEqual(result.created, 6)
        self.assertEqual(result.skipped, 0)
        payment = FinancialTransaction.objects.get(
            external_id=_external_id(date(2026, 9, 2), "GEL", Decimal("-5.00"), MERCHANT_DETAILS)
        )
        self.assertEqual(payment.account, gel)
        self.assertEqual(payment.date, date(2026, 9, 2))
        self.assertEqual(payment.amount, Decimal("-5.00"))
        self.assertEqual(payment.kind, FinancialTransaction.KIND_EXPENSE)
        self.assertEqual(payment.income_type, FinancialTransaction.INCOME_TRANSFER)
        self.assertEqual(payment.counterparty_name, "Batumi Airport, Batumi, Airport Highway")
        self.assertEqual(payment.counterparty_iban, "")
        self.assertEqual(payment.description, MERCHANT_DETAILS.strip())
        income = FinancialTransaction.objects.get(
            external_id=_external_id(
                date(2026, 9, 2),
                "GEL",
                Decimal("51.70"),
                INCOME_CONVERSION_DETAILS,
            )
        )
        self.assertEqual(income.account, gel)
        self.assertEqual(income.amount, Decimal("51.70"))
        self.assertEqual(income.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(income.income_type, FinancialTransaction.INCOME_TRANSFER)
        self.assertEqual(income.counterparty_name, "")
        self.assertEqual(income.counterparty_iban, "")
        conversion = FinancialTransaction.objects.get(
            external_id=_external_id(
                date(2026, 9, 2),
                "USD",
                Decimal("-20.17"),
                EXPENSE_CONVERSION_DETAILS,
            )
        )
        self.assertEqual(conversion.account, usd)
        self.assertEqual(conversion.amount, Decimal("-20.17"))
        self.assertEqual(conversion.kind, FinancialTransaction.KIND_EXPENSE)
        outgoing = FinancialTransaction.objects.get(
            external_id=_external_id(date(2026, 9, 6), "GEL", Decimal("-200.00"), OUTGOING_DETAILS)
        )
        self.assertEqual(outgoing.counterparty_name, "im dzmitri haravets")
        self.assertEqual(outgoing.counterparty_iban, "GE13BG0000000609702806GEL")
        self.assertEqual(outgoing.kind, FinancialTransaction.KIND_EXPENSE)
        incoming = FinancialTransaction.objects.get(
            external_id=_external_id(date(2026, 9, 8), "GEL", Decimal("1350.00"), INCOMING_DETAILS)
        )
        self.assertEqual(incoming.counterparty_name, "i/m daria kharlamova")
        self.assertEqual(incoming.counterparty_iban, "GE07TB7154245068100036")
        self.assertEqual(incoming.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(incoming.amount, Decimal("1350.00"))
        withdrawal = FinancialTransaction.objects.get(
            external_id=_external_id(date(2026, 9, 11), "GEL", Decimal("-1000.00"), WITHDRAWAL_DETAILS)
        )
        self.assertEqual(withdrawal.counterparty_name, "Bank of Georgia, Batumi, Kobaladze St. 15/73 (2)")
        self.assertEqual(withdrawal.counterparty_iban, "")
        self.assertEqual(withdrawal.kind, FinancialTransaction.KIND_EXPENSE)
        self.assertEqual(withdrawal.amount, Decimal("-1000.00"))

    def test_skips_duplicates(self):
        _account(GEL_IBAN, "GEL")
        _account(USD_IBAN, "USD")
        statement = _statement(_rows())
        import_bog_personal_statement(statement)
        statement.seek(0)
        again = import_bog_personal_statement(statement)
        self.assertEqual(again.error, "")
        self.assertEqual(again.created, 0)
        self.assertEqual(again.skipped, 6)
        self.assertEqual(FinancialTransaction.objects.count(), 6)

    def test_applies_expense_rule(self):
        _account(GEL_IBAN, "GEL")
        rule = FinancialTransactionRule.objects.create(
            name="Airport",
            expense_type=FinancialTransaction.EXPENSE_TAXI,
            priority=1,
            operator=FinancialTransactionRule.OPERATOR_OR,
            is_active=True,
        )
        FinancialTransactionRuleCondition.objects.create(
            rule=rule,
            field=FinancialTransactionRuleCondition.FIELD_DESCRIPTION,
            pattern="*Batumi Airport*",
        )
        result = import_bog_personal_statement(
            _statement([{"Date": "02/09/2026", "Details": MERCHANT_DETAILS, "GEL": -5}])
        )
        self.assertEqual(result.error, "")
        payment = FinancialTransaction.objects.get()
        self.assertEqual(payment.expense_type, FinancialTransaction.EXPENSE_TAXI)

    def test_unknown_iban_creates_nothing(self):
        _account(GEL_IBAN, "GEL")
        result = import_bog_personal_statement(
            _statement(
                [
                    {"Date": "02/09/2026", "Details": MERCHANT_DETAILS, "GEL": -5},
                    {"Date": "02/09/2026", "Details": EXPENSE_CONVERSION_DETAILS, "USD": -20.17},
                ]
            )
        )
        self.assertIn(USD_IBAN, result.error)
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)
