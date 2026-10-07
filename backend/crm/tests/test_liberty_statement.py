from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook

from crm.liberty_statement import import_liberty_statement
from crm.models import FinancialAccount, FinancialTransaction

IBAN = "GE53LB0112183991465000"
EXCEL_EPOCH = datetime(1899, 12, 30)
HEADERS = [
    "Date",
    "Description",
    "Paid Out",
    "Paid In",
    "Document Number",
    "Partner's Name",
    "Partner's Account",
]
ONLINE_DESCRIPTION = (
    "Reimbursement/ ანგარიშსწორება: ORDERID 4799853,  Georgia 08/07/2026 23:03:35"
    " | AUTH. No 836752 | GEL 9.00 - 9001166 | Terminal No IA01166"
)
TERMINAL_DESCRIPTION = "ბარათით განაღდებები 8.06.2026"
TRANSFER_DESCRIPTION = "Dividends payment"
FEE_DESCRIPTION = "Wire Transfer Fee / თანხის გადარიცხვის საკომისიო"
SHARED_DOCUMENT_NUMBER = "260701102439063"


def _excel_serial(value: datetime) -> int:
    return (value - EXCEL_EPOCH).days


def _statement(rows: list[dict]) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["B5"] = "Account No:"
    summary["C5"] = IBAN
    worksheet = workbook.create_sheet("account_statement")
    worksheet["A1"] = "თარიღი"
    for index, header in enumerate(HEADERS, start=1):
        worksheet.cell(2, index, header)
    for offset, row in enumerate(rows):
        excel_row = 3 + offset
        for index, header in enumerate(HEADERS, start=1):
            cell = worksheet.cell(excel_row, index, row[header])
            if header == "Date":
                cell.number_format = "mm-dd-yy"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _account() -> FinancialAccount:
    return FinancialAccount.objects.create(
        name="Liberty",
        kind=FinancialAccount.KIND_BANK,
        iban=IBAN,
        currency="GEL",
    )


def _online() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 7, 9)),
        "Description": ONLINE_DESCRIPTION,
        "Paid Out": None,
        "Paid In": 8.82,
        "Document Number": 576185545,
        "Partner's Name": "SWEET&CHILL - სატრანზიტო",
        "Partner's Account": "GE34LB0010145014881166",
    }


def _terminal() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 6, 9)),
        "Description": TERMINAL_DESCRIPTION,
        "Paid Out": None,
        "Paid In": 9.8,
        "Document Number": 1709,
        "Partner's Name": "ინდივიდუალური მეწარმე ფილიპპ ბაკანოვ",
        "Partner's Account": "GE93LB0100450106000614",
    }


def _transfer() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 7, 1)),
        "Description": TRANSFER_DESCRIPTION,
        "Paid Out": 550,
        "Paid In": None,
        "Document Number": SHARED_DOCUMENT_NUMBER,
        "Partner's Name": "FILIPP BAKANOV",
        "Partner's Account": "GE14BG0000000537815409",
    }


def _fee() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 7, 1)),
        "Description": FEE_DESCRIPTION,
        "Paid Out": 2,
        "Paid In": None,
        "Document Number": SHARED_DOCUMENT_NUMBER,
        "Partner's Name": "შემოსავალი ერთჯერადი გადარიცხვებიდან - ლარი",
        "Partner's Account": "GE67LB0016864040710012",
    }


class LibertyStatementImportTests(TestCase):
    def test_creates_online_terminal_and_shared_document_number(self):
        account = _account()
        result = import_liberty_statement(_statement([_online(), _terminal(), _transfer(), _fee()]))
        self.assertEqual(result.error, "")
        self.assertEqual(result.created, 4)
        self.assertEqual(result.skipped, 0)
        online = FinancialTransaction.objects.get(external_id=f"576185545|{ONLINE_DESCRIPTION}")
        self.assertEqual(online.account, account)
        self.assertEqual(online.date, datetime(2026, 7, 9).date())
        self.assertEqual(online.amount, Decimal("8.82"))
        self.assertEqual(online.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(online.income_type, FinancialTransaction.INCOME_ONLINE)
        self.assertEqual(online.counterparty_name, "SWEET&CHILL - სატრანზიტო")
        self.assertEqual(online.counterparty_iban, "GE34LB0010145014881166")
        self.assertEqual(online.description, ONLINE_DESCRIPTION)
        terminal = FinancialTransaction.objects.get(external_id=f"1709|{TERMINAL_DESCRIPTION}")
        self.assertEqual(terminal.amount, Decimal("9.80"))
        self.assertEqual(terminal.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(terminal.income_type, FinancialTransaction.INCOME_TERMINAL)
        transfer = FinancialTransaction.objects.get(
            external_id=f"{SHARED_DOCUMENT_NUMBER}|{TRANSFER_DESCRIPTION}"
        )
        fee = FinancialTransaction.objects.get(external_id=f"{SHARED_DOCUMENT_NUMBER}|{FEE_DESCRIPTION}")
        self.assertEqual(transfer.amount, Decimal("-550.00"))
        self.assertEqual(transfer.kind, FinancialTransaction.KIND_EXPENSE)
        self.assertEqual(transfer.income_type, "")
        self.assertEqual(fee.amount, Decimal("-2.00"))
        self.assertEqual(fee.kind, FinancialTransaction.KIND_EXPENSE)
        self.assertEqual(fee.income_type, "")

    def test_skips_duplicates(self):
        _account()
        statement = _statement([_online(), _terminal(), _transfer(), _fee()])
        import_liberty_statement(statement)
        statement.seek(0)
        again = import_liberty_statement(statement)
        self.assertEqual(again.error, "")
        self.assertEqual(again.created, 0)
        self.assertEqual(again.skipped, 4)
        self.assertEqual(FinancialTransaction.objects.count(), 4)

    def test_unknown_iban_creates_nothing(self):
        result = import_liberty_statement(_statement([_online(), _terminal()]))
        self.assertIn(IBAN, result.error)
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)
