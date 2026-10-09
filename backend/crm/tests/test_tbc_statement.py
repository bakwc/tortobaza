from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook

from crm.models import (
    FinancialAccount,
    FinancialTransaction,
    FinancialTransactionRule,
    FinancialTransactionRuleCondition,
)
from crm.tbc_statement import import_tbc_statement

HEADERS = [
    "Дата",
    "Описание операции",
    "Дополнительная информация",
    "Списание, GEL",
    "Поступление, GEL",
    "Остаток, GEL",
    "Стр. PDF",
]
INCOME_OPERATION = "Private transfer within TBC"
INCOME_ADDITIONAL = "ავთანდილ გორგაძე, TBCBGE22, GE29TB7649345064300130"
EXPENSE_OPERATION = "POS - Vip Pay*YANDEX.GO, 4.70 GEL, Sep 1 2026 8:26PM"
EXPENSE_ADDITIONAL = "TBCBank merchant, TBCBGE22, GE35TB0002511341111111"
INCOME_EXTERNAL_ID = (
    "2026-09-02|Private transfer within TBC|"
    "ავთანდილ გორგაძე, TBCBGE22, GE29TB7649345064300130|90.00"
)
EXPENSE_EXTERNAL_ID = (
    "2026-09-02|POS - Vip Pay*YANDEX.GO, 4.70 GEL, Sep 1 2026 8:26PM|"
    "TBCBank merchant, TBCBGE22, GE35TB0002511341111111|-4.70"
)


def _statement(rows: list[dict]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Операции"
    for index, header in enumerate(HEADERS, start=1):
        worksheet.cell(1, index, header)
    for offset, row in enumerate(rows):
        excel_row = 2 + offset
        for index, header in enumerate(HEADERS, start=1):
            worksheet.cell(excel_row, index, row[header])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _account(name: str, bank_name: str) -> FinancialAccount:
    return FinancialAccount.objects.create(
        name=name,
        kind=FinancialAccount.KIND_BANK,
        bank_name=bank_name,
        currency="GEL",
    )


def _income() -> dict:
    return {
        "Дата": datetime(2026, 9, 2),
        "Описание операции": INCOME_OPERATION,
        "Дополнительная информация": INCOME_ADDITIONAL,
        "Списание, GEL": None,
        "Поступление, GEL": 90,
        "Остаток, GEL": 368.93,
        "Стр. PDF": 1,
    }


def _expense() -> dict:
    return {
        "Дата": datetime(2026, 9, 2),
        "Описание операции": EXPENSE_OPERATION,
        "Дополнительная информация": EXPENSE_ADDITIONAL,
        "Списание, GEL": 4.7,
        "Поступление, GEL": None,
        "Остаток, GEL": 358.13,
        "Стр. PDF": 2,
    }


class TbcStatementImportTests(TestCase):
    def test_creates_income_and_expense(self):
        account = _account("TBC card", "TBC")
        result = import_tbc_statement(_statement([_income(), _expense()]))
        self.assertEqual(result.error, "")
        self.assertEqual(result.created, 2)
        self.assertEqual(result.skipped, 0)
        income = FinancialTransaction.objects.get(external_id=INCOME_EXTERNAL_ID)
        self.assertEqual(income.account, account)
        self.assertEqual(income.date, datetime(2026, 9, 2).date())
        self.assertEqual(income.amount, Decimal("90.00"))
        self.assertEqual(income.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(income.income_type, FinancialTransaction.INCOME_TRANSFER)
        self.assertEqual(income.counterparty_name, "ავთანდილ გორგაძე")
        self.assertEqual(income.counterparty_iban, "GE29TB7649345064300130")
        self.assertEqual(income.description, f"{INCOME_OPERATION} | {INCOME_ADDITIONAL}")
        expense = FinancialTransaction.objects.get(external_id=EXPENSE_EXTERNAL_ID)
        self.assertEqual(expense.account, account)
        self.assertEqual(expense.amount, Decimal("-4.70"))
        self.assertEqual(expense.kind, FinancialTransaction.KIND_EXPENSE)
        self.assertEqual(expense.income_type, "")
        self.assertEqual(expense.counterparty_name, "TBCBank merchant")
        self.assertEqual(expense.counterparty_iban, "GE35TB0002511341111111")
        self.assertEqual(expense.description, f"{EXPENSE_OPERATION} | {EXPENSE_ADDITIONAL}")

    def test_skips_duplicates(self):
        _account("TBC card", "TBC")
        statement = _statement([_income(), _expense()])
        import_tbc_statement(statement)
        statement.seek(0)
        again = import_tbc_statement(statement)
        self.assertEqual(again.error, "")
        self.assertEqual(again.created, 0)
        self.assertEqual(again.skipped, 2)
        self.assertEqual(FinancialTransaction.objects.count(), 2)
        fresh = _expense()
        fresh["Списание, GEL"] = 18
        mixed = import_tbc_statement(_statement([_income(), _income(), fresh]))
        self.assertEqual(mixed.error, "")
        self.assertEqual(mixed.created, 1)
        self.assertEqual(mixed.skipped, 2)
        self.assertEqual(FinancialTransaction.objects.count(), 3)

    def test_applies_expense_rule(self):
        _account("TBC card", "TBC")
        rule = FinancialTransactionRule.objects.create(
            name="Yandex",
            expense_type=FinancialTransaction.EXPENSE_FEES,
            priority=1,
            operator=FinancialTransactionRule.OPERATOR_OR,
            is_active=True,
        )
        FinancialTransactionRuleCondition.objects.create(
            rule=rule,
            field=FinancialTransactionRuleCondition.FIELD_DESCRIPTION,
            pattern="*YANDEX.GO*",
        )
        result = import_tbc_statement(_statement([_income(), _expense()]))
        self.assertEqual(result.error, "")
        expense = FinancialTransaction.objects.get(external_id=EXPENSE_EXTERNAL_ID)
        income = FinancialTransaction.objects.get(external_id=INCOME_EXTERNAL_ID)
        self.assertEqual(expense.expense_type, FinancialTransaction.EXPENSE_FEES)
        self.assertEqual(income.expense_type, "")

    def test_missing_account_creates_nothing(self):
        _account("TBC", "")
        result = import_tbc_statement(_statement([_income(), _expense()]))
        self.assertEqual(result.error, "TBC account not found")
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)

    def test_ambiguous_account_creates_nothing(self):
        _account("TBC card", "TBC")
        _account("TBC business", "TBC")
        result = import_tbc_statement(_statement([_income()]))
        self.assertEqual(result.error, "TBC account matched more than once")
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)

    def test_both_amounts_creates_nothing(self):
        _account("TBC card", "TBC")
        row = _income()
        row["Списание, GEL"] = 10
        result = import_tbc_statement(_statement([row, _expense()]))
        self.assertEqual(result.error, "Row 2: both debit and credit are set")
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)

    def test_empty_amounts_creates_nothing(self):
        _account("TBC card", "TBC")
        row = _income()
        row["Поступление, GEL"] = None
        result = import_tbc_statement(_statement([row, _expense()]))
        self.assertEqual(result.error, "Row 2: debit and credit are empty")
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)

    def test_invalid_additional_information_creates_nothing(self):
        _account("TBC card", "TBC")
        row = _income()
        row["Дополнительная информация"] = "only a name"
        result = import_tbc_statement(_statement([row, _expense()]))
        self.assertEqual(result.error, "Row 2: additional information must have three parts")
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)

    def test_missing_iban_creates_nothing(self):
        _account("TBC card", "TBC")
        row = _income()
        row["Дополнительная информация"] = "ავთანდილ გორგაძე, TBCBGE22,"
        result = import_tbc_statement(_statement([row]))
        self.assertEqual(result.error, "Row 2: missing IBAN")
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)
