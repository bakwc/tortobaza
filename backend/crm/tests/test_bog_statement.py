from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook

from crm.bog_statement import import_bog_statement
from crm.models import FinancialAccount, FinancialTransaction

GEL_IBAN = "GE94BG0000000612361573GEL"
USD_IBAN = "GE94BG0000000612361573USD"
EXCEL_EPOCH = datetime(1899, 12, 30)
HEADERS = [
    "Date",
    "Account N",
    "Debit",
    "Credit",
    "Entry Comment",
    "Operation ID",
    "Sender Name",
    "Sender Account N",
    " Recipient Name",
    "Recipient Account N",
    "Nomination",
    "Amount",
]


def _excel_serial(value: datetime) -> int:
    return (value - EXCEL_EPOCH).days


def _statement(rows: list[dict]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Statement of Account"
    worksheet["A4"] = "Statement"
    for index, header in enumerate(HEADERS, start=1):
        worksheet.cell(9, index, header)
    for offset, row in enumerate(rows):
        excel_row = 10 + offset
        for index, header in enumerate(HEADERS, start=1):
            cell = worksheet.cell(excel_row, index, row[header.strip()])
            if header == "Date":
                cell.number_format = "mm-dd-yy"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _account(iban: str, currency: str, name: str) -> FinancialAccount:
    return FinancialAccount.objects.create(
        name=name,
        kind=FinancialAccount.KIND_BANK,
        iban=iban,
        currency=currency,
    )


def _income() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 6, 12)),
        "Account N": GEL_IBAN,
        "Debit": None,
        "Credit": 125.97,
        "Entry Comment": "Incoming Transfer",
        "Operation ID": 118945436567,
        "Sender Name": "shps volt jorjia",
        "Sender Account N": "GE62BG0000000100816256GEL",
        "Recipient Name": "INDIVIDUAL ENTREPRENEUR FILIPP BAKANOV",
        "Recipient Account N": GEL_IBAN,
        "Nomination": "Wolt payout",
        "Amount": 125.97,
    }


def _expense() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 7, 10)),
        "Account N": GEL_IBAN,
        "Debit": 5,
        "Credit": None,
        "Entry Comment": "Business Package S Maintenance Fee ",
        "Operation ID": 119001305817,
        "Sender Name": "INDIVIDUAL ENTREPRENEUR FILIPP BAKANOV",
        "Sender Account N": GEL_IBAN,
        "Recipient Name": None,
        "Recipient Account N": None,
        "Nomination": "",
        "Amount": -5,
    }


def _usd_income() -> dict:
    return {
        "Date": _excel_serial(datetime(2026, 9, 1)),
        "Account N": USD_IBAN,
        "Debit": None,
        "Credit": 26.93,
        "Entry Comment": "Incoming Transfer",
        "Operation ID": 124553319539,
        "Sender Name": "HRYVNIUK TETIANA BOHDANIVNA",
        "Sender Account N": "UA693052990000026209896773055",
        "Recipient Name": "INDIVIDUAL ENTREPRENEUR FILIPP BAKANOV",
        "Recipient Account N": USD_IBAN,
        "Nomination": "PRPMNT OF CAKE",
        "Amount": 26.93,
    }


class BogStatementImportTests(TestCase):
    def test_creates_income_and_expense(self):
        gel = _account(GEL_IBAN, "GEL", "SWEET CHILL")
        usd = _account(USD_IBAN, "USD", "SWEET CHILL USD")
        result = import_bog_statement(_statement([_income(), _expense(), _usd_income()]))
        self.assertEqual(result.error, "")
        self.assertEqual(result.created, 3)
        self.assertEqual(result.skipped, 0)
        income = FinancialTransaction.objects.get(external_id="118945436567")
        self.assertEqual(income.account, gel)
        self.assertEqual(income.date, datetime(2026, 6, 12).date())
        self.assertEqual(income.amount, Decimal("125.97"))
        self.assertEqual(income.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(income.counterparty_name, "shps volt jorjia")
        self.assertEqual(income.counterparty_iban, "GE62BG0000000100816256GEL")
        self.assertEqual(income.description, "Wolt payout")
        expense = FinancialTransaction.objects.get(external_id="119001305817")
        self.assertEqual(expense.account, gel)
        self.assertEqual(expense.amount, Decimal("-5.00"))
        self.assertEqual(expense.kind, FinancialTransaction.KIND_EXPENSE)
        self.assertEqual(expense.counterparty_name, "")
        self.assertEqual(expense.counterparty_iban, "")
        self.assertEqual(expense.description, "Business Package S Maintenance Fee")
        usd_income = FinancialTransaction.objects.get(external_id="124553319539")
        self.assertEqual(usd_income.account, usd)
        self.assertEqual(usd_income.kind, FinancialTransaction.KIND_INCOME)
        self.assertEqual(usd_income.counterparty_iban, "UA693052990000026209896773055")

    def test_skips_duplicates(self):
        _account(GEL_IBAN, "GEL", "SWEET CHILL")
        _account(USD_IBAN, "USD", "SWEET CHILL USD")
        statement = _statement([_income(), _expense(), _usd_income()])
        import_bog_statement(statement)
        statement.seek(0)
        again = import_bog_statement(statement)
        self.assertEqual(again.error, "")
        self.assertEqual(again.created, 0)
        self.assertEqual(again.skipped, 3)
        self.assertEqual(FinancialTransaction.objects.count(), 3)
        duplicate = _income()
        fresh = _expense()
        fresh["Operation ID"] = 119201857423
        fresh["Amount"] = -18
        fresh["Debit"] = 18
        mixed = import_bog_statement(_statement([duplicate, duplicate, fresh]))
        self.assertEqual(mixed.error, "")
        self.assertEqual(mixed.created, 1)
        self.assertEqual(mixed.skipped, 2)
        self.assertEqual(FinancialTransaction.objects.count(), 4)

    def test_unknown_iban_creates_nothing(self):
        _account(GEL_IBAN, "GEL", "SWEET CHILL")
        result = import_bog_statement(_statement([_income(), _usd_income()]))
        self.assertIn(USD_IBAN, result.error)
        self.assertEqual(result.created, 0)
        self.assertEqual(FinancialTransaction.objects.count(), 0)
