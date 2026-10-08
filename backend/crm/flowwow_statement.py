import calendar
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from django.db import transaction
from openpyxl import load_workbook

from crm.bog_statement import column_value, decimal_amount, is_filled, text_value
from crm.models import FinancialAccount, FinancialTransaction

ACCOUNT_NAME = "Flowwow"
COUNTERPARTY_NAME = "Flowwow"
SHOP_CURRENCY = "GEL"
REQUIRED_COLUMNS = (
    "Contract number",
    "Contract date",
    "Order number",
    "Order payment date",
    "Order delivery date",
    "Transaction type",
    "Transaction amount",
    "Transaction currency",
    "Transaction rate",
    "Transaction amount in shop currency",
)
INCOME_TYPES = {
    "Paid by customer",
    "Delivery surcharge",
    "Compensation for applied discount",
}
EXPENSE_TYPES = {
    "Card processing fee",
    "Bonus points",
    "Fee",
}
WITHDRAWAL_TYPE = "Withdrawal"
PAYMENT_DATE_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4}) (\d{2}):(\d{2}):(\d{2})$")


@dataclass
class ParsedFlowwowRow:
    date: date
    amount: Decimal
    kind: str
    counterparty_name: str
    description: str
    external_id: str


@dataclass
class FlowwowStatementImportResult:
    created: int
    skipped: int
    error: str


def import_flowwow_statement(file) -> FlowwowStatementImportResult:
    rows, error = parse_flowwow_statement(file)
    if error:
        return FlowwowStatementImportResult(created=0, skipped=0, error=error)
    if not rows:
        return FlowwowStatementImportResult(created=0, skipped=0, error="")

    accounts = list(FinancialAccount.objects.filter(name=ACCOUNT_NAME))
    if len(accounts) == 0:
        return FlowwowStatementImportResult(created=0, skipped=0, error="Flowwow account not found")
    if len(accounts) > 1:
        return FlowwowStatementImportResult(
            created=0,
            skipped=0,
            error="Flowwow account matched more than once",
        )
    account = accounts[0]

    external_ids = [row.external_id for row in rows]
    with transaction.atomic():
        existing = set(
            FinancialTransaction.objects.filter(
                account=account,
                external_id__in=external_ids,
            ).values_list("external_id", flat=True)
        )
        to_create: list[FinancialTransaction] = []
        skipped = 0
        for row in rows:
            if row.external_id in existing:
                skipped += 1
                continue
            existing.add(row.external_id)
            to_create.append(
                FinancialTransaction(
                    account=account,
                    date=row.date,
                    amount=row.amount,
                    kind=row.kind,
                    counterparty_name=row.counterparty_name,
                    description=row.description,
                    external_id=row.external_id,
                )
            )
        FinancialTransaction.objects.bulk_create(to_create)
    return FlowwowStatementImportResult(created=len(to_create), skipped=skipped, error="")


def parse_flowwow_statement(file) -> tuple[list[ParsedFlowwowRow], str]:
    workbook = load_workbook(file, data_only=False)
    worksheet = workbook.worksheets[0]
    header_map: dict[str, int] | None = None
    rows: list[ParsedFlowwowRow] = []
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        if header_map is None:
            names: dict[str, int] = {}
            for index, value in enumerate(values):
                if isinstance(value, str) and value.strip():
                    names[value.strip()] = index
            if "Order number" in names and "Transaction type" in names:
                missing_columns = [name for name in REQUIRED_COLUMNS if name not in names]
                if missing_columns:
                    return [], f"Missing columns: {', '.join(missing_columns)}"
                header_map = names
            continue
        if not any(is_filled(value) for value in values):
            continue
        row_number = excel_row[0].row
        parsed, error = parse_flowwow_row(values, header_map, row_number)
        if error:
            return [], error
        rows.append(parsed)
    if header_map is None:
        return [], "Statement header not found"
    return rows, ""


def parse_flowwow_row(
    values: list,
    header_map: dict[str, int],
    row_number: int,
) -> tuple[ParsedFlowwowRow | None, str]:
    transaction_type = text_value(column_value(values, header_map, "Transaction type"))
    if transaction_type == "":
        return None, f"Row {row_number}: missing Transaction type"
    order_number = text_value(column_value(values, header_map, "Order number"))
    if order_number == "":
        return None, f"Row {row_number}: missing Order number"
    raw_shop_amount = column_value(values, header_map, "Transaction amount in shop currency")
    if not is_filled(raw_shop_amount):
        return None, f"Row {row_number}: missing Transaction amount in shop currency"
    amount = decimal_amount(raw_shop_amount)
    currency = text_value(column_value(values, header_map, "Transaction currency"))
    if currency == "":
        return None, f"Row {row_number}: missing Transaction currency"
    payment = parse_payment_datetime(column_value(values, header_map, "Order payment date"))
    if payment is None:
        return None, f"Row {row_number}: invalid date"
    kind, error = transaction_kind(transaction_type, amount, row_number)
    if error:
        return None, error
    if currency == SHOP_CURRENCY:
        description = transaction_type
    else:
        raw_amount = column_value(values, header_map, "Transaction amount")
        if not is_filled(raw_amount):
            return None, f"Row {row_number}: missing Transaction amount"
        description = f"{transaction_type} {decimal_amount(raw_amount)} {currency}"
    external_id = (
        f"{order_number}|{transaction_type}|{payment.strftime('%Y-%m-%d %H:%M:%S')}|{amount}"
    )
    return (
        ParsedFlowwowRow(
            date=payment.date(),
            amount=amount,
            kind=kind,
            counterparty_name=COUNTERPARTY_NAME,
            description=description,
            external_id=external_id,
        ),
        "",
    )


def transaction_kind(transaction_type: str, amount: Decimal, row_number: int) -> tuple[str, str]:
    if transaction_type in INCOME_TYPES:
        if amount <= 0:
            return "", f"Row {row_number}: amount sign does not match {transaction_type}"
        return FinancialTransaction.KIND_INCOME, ""
    if transaction_type in EXPENSE_TYPES:
        if amount >= 0:
            return "", f"Row {row_number}: amount sign does not match {transaction_type}"
        return FinancialTransaction.KIND_EXPENSE, ""
    if transaction_type == WITHDRAWAL_TYPE:
        if amount > 0:
            return "", f"Row {row_number}: amount sign does not match {transaction_type}"
        return FinancialTransaction.KIND_WITHDRAWAL, ""
    return "", f"Row {row_number}: unknown transaction type"


def parse_payment_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if not isinstance(value, str):
        return None
    match = PAYMENT_DATE_RE.match(value.strip())
    if match is None:
        return None
    day = int(match.group(1))
    month = int(match.group(2))
    year = int(match.group(3))
    hour = int(match.group(4))
    minute = int(match.group(5))
    second = int(match.group(6))
    if month < 1 or month > 12:
        return None
    if day < 1 or day > calendar.monthrange(year, month)[1]:
        return None
    if hour > 23 or minute > 59 or second > 59:
        return None
    return datetime(year, month, day, hour, minute, second)
