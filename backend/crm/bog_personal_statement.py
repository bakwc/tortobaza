import calendar
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from crm.bog_statement import column_value, decimal_amount, is_filled, parse_date, text_value
from crm.models import FinancialAccount, FinancialTransaction
from crm.transaction_rules import assign_expense_types
from tortobaza.finance_db import current_finance_db

DETAILS_SHEET = "Details"
TRANSACTIONS_SHEET = "Transactions"
ACCOUNT_LABEL = "Account No:"
CURRENCIES = ("GEL", "USD", "EUR", "GBP")
REQUIRED_COLUMNS = ("Date", "Details", *CURRENCIES)
COUNTERPARTY_NAME_KEYS = ("Beneficiary", "Sender", "Merchant", "ATM")


@dataclass
class ParsedBogPersonalRow:
    account_iban: str
    date: date
    amount: Decimal
    kind: str
    income_type: str
    counterparty_name: str
    counterparty_iban: str
    description: str
    external_id: str


@dataclass
class BogPersonalStatementImportResult:
    created: int
    skipped: int
    error: str


def import_bog_personal_statement(file) -> BogPersonalStatementImportResult:
    rows, error = parse_bog_personal_statement(file)
    if error:
        return BogPersonalStatementImportResult(created=0, skipped=0, error=error)

    ibans = {row.account_iban for row in rows}
    found = list(FinancialAccount.objects.filter(iban__in=ibans))
    accounts: dict[str, FinancialAccount] = {}
    ambiguous: list[str] = []
    for account in found:
        if account.iban in accounts:
            ambiguous.append(account.iban)
        else:
            accounts[account.iban] = account
    missing = sorted(ibans - set(accounts))
    ambiguous = sorted(set(ambiguous))
    messages: list[str] = []
    if missing:
        messages.append(_("Unknown accounts: %(accounts)s") % {"accounts": ", ".join(missing)})
    if ambiguous:
        messages.append(_("Accounts matched more than once: %(accounts)s") % {"accounts": ", ".join(ambiguous)})
    if messages:
        return BogPersonalStatementImportResult(created=0, skipped=0, error="; ".join(messages))

    external_ids = [row.external_id for row in rows]
    account_ids = [account.id for account in accounts.values()]
    with transaction.atomic(using=current_finance_db()):
        existing = set(
            FinancialTransaction.objects.filter(
                account_id__in=account_ids,
                external_id__in=external_ids,
            ).values_list("account_id", "external_id")
        )
        to_create: list[FinancialTransaction] = []
        skipped = 0
        for row in rows:
            account = accounts[row.account_iban]
            key = (account.id, row.external_id)
            if key in existing:
                skipped += 1
                continue
            existing.add(key)
            to_create.append(
                FinancialTransaction(
                    account=account,
                    date=row.date,
                    amount=row.amount,
                    kind=row.kind,
                    income_type=row.income_type,
                    counterparty_name=row.counterparty_name,
                    counterparty_iban=row.counterparty_iban,
                    description=row.description,
                    external_id=row.external_id,
                )
            )
        assign_expense_types(to_create)
        FinancialTransaction.objects.bulk_create(to_create)
    return BogPersonalStatementImportResult(created=len(to_create), skipped=skipped, error="")


def parse_bog_personal_statement(file) -> tuple[list[ParsedBogPersonalRow], str]:
    workbook = load_workbook(file, data_only=False)
    if DETAILS_SHEET not in workbook.sheetnames:
        return [], _("Sheet 'Details' not found")
    if TRANSACTIONS_SHEET not in workbook.sheetnames:
        return [], _("Sheet 'Transactions' not found")
    account_number, error = read_account_number(workbook[DETAILS_SHEET])
    if error:
        return [], error
    worksheet = workbook[TRANSACTIONS_SHEET]
    header_map: dict[str, int] | None = None
    rows: list[ParsedBogPersonalRow] = []
    seen_external_ids: set[str] = set()
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        if header_map is None:
            names: dict[str, int] = {}
            for index, value in enumerate(values):
                if isinstance(value, str) and value.strip():
                    names[value.strip()] = index
            if "Date" in names and "Details" in names:
                missing_columns = [name for name in REQUIRED_COLUMNS if name not in names]
                if missing_columns:
                    return [], _("Missing columns: %(columns)s") % {"columns": ", ".join(missing_columns)}
                header_map = names
            continue
        if not any(is_filled(value) for value in values):
            continue
        row_number = excel_row[0].row
        parsed, error = parse_bog_personal_row(values, header_map, row_number, account_number)
        if error:
            return [], error
        if parsed.external_id in seen_external_ids:
            return [], _("Row %(row_number)s: duplicate external id") % {"row_number": row_number}
        seen_external_ids.add(parsed.external_id)
        rows.append(parsed)
    if header_map is None:
        return [], _("Statement header not found")
    return rows, ""


def parse_bog_personal_row(
    values: list,
    header_map: dict[str, int],
    row_number: int,
    account_number: str,
) -> tuple[ParsedBogPersonalRow | None, str]:
    description = text_value(column_value(values, header_map, "Details"))
    if description == "":
        return None, _("Row %(row_number)s: missing Details") % {"row_number": row_number}
    filled: list[tuple[str, object]] = []
    for currency in CURRENCIES:
        raw_amount = column_value(values, header_map, currency)
        if is_filled(raw_amount):
            filled.append((currency, raw_amount))
    if len(filled) == 0:
        return None, _("Row %(row_number)s: missing amount") % {"row_number": row_number}
    if len(filled) > 1:
        return None, _("Row %(row_number)s: multiple currencies") % {"row_number": row_number}
    currency, raw_amount = filled[0]
    amount = decimal_amount(raw_amount)
    if amount == 0:
        return None, _("Row %(row_number)s: amount is zero") % {"row_number": row_number}
    if amount > 0:
        kind = FinancialTransaction.KIND_INCOME
    else:
        kind = FinancialTransaction.KIND_EXPENSE
    parsed_date = parse_statement_date(column_value(values, header_map, "Date"))
    if parsed_date is None:
        return None, _("Row %(row_number)s: invalid date") % {"row_number": row_number}
    external_id = f"{parsed_date.isoformat()}|{currency}|{amount}|{description}"
    max_length = FinancialTransaction._meta.get_field("external_id").max_length
    if len(external_id) > max_length:
        return None, _("Row %(row_number)s: external id is too long") % {"row_number": row_number}
    counterparty_name, counterparty_iban = counterparty_from_details(description)
    return (
        ParsedBogPersonalRow(
            account_iban=f"{account_number}{currency}",
            date=parsed_date,
            amount=amount,
            kind=kind,
            income_type=FinancialTransaction.INCOME_TRANSFER,
            counterparty_name=counterparty_name,
            counterparty_iban=counterparty_iban,
            description=description,
            external_id=external_id,
        ),
        "",
    )


def counterparty_from_details(details: str) -> tuple[str, str]:
    fields: dict[str, str] = {}
    for segment in details.split(";"):
        key, separator, value = segment.strip().partition(":")
        if separator == "":
            continue
        name = key.strip()
        if name not in fields:
            fields[name] = value.strip()
    counterparty_name = ""
    for key in COUNTERPARTY_NAME_KEYS:
        if key in fields and fields[key] != "":
            counterparty_name = fields[key]
            break
    if "Account" in fields:
        counterparty_iban = fields["Account"]
    else:
        counterparty_iban = ""
    return counterparty_name, counterparty_iban


def parse_statement_date(value) -> date | None:
    if isinstance(value, str):
        parts = value.strip().split("/")
        if len(parts) != 3:
            return None
        day_text, month_text, year_text = parts
        if len(day_text) != 2 or len(month_text) != 2 or len(year_text) != 4:
            return None
        if not day_text.isdigit() or not month_text.isdigit() or not year_text.isdigit():
            return None
        day = int(day_text)
        month = int(month_text)
        year = int(year_text)
        if month < 1 or month > 12:
            return None
        if day < 1 or day > calendar.monthrange(year, month)[1]:
            return None
        return date(year, month, day)
    return parse_date(value)


def read_account_number(worksheet) -> tuple[str, str]:
    found: list[str] = []
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        for index, value in enumerate(values):
            if isinstance(value, str) and value.strip() == ACCOUNT_LABEL:
                number = ""
                for next_value in values[index + 1 :]:
                    if is_filled(next_value):
                        number = text_value(next_value)
                        break
                found.append(number)
    numbers = sorted({number for number in found if number})
    if len(numbers) == 1 and "" not in found:
        return numbers[0], ""
    if len(numbers) > 1:
        return "", _("Multiple account numbers: %(numbers)s") % {"numbers": ", ".join(numbers)}
    return "", _("Account No not found")
