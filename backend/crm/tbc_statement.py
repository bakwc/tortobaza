from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from crm.bog_statement import column_value, decimal_amount, is_filled, parse_date, text_value
from crm.models import FinancialAccount, FinancialTransaction
from crm.transaction_rules import assign_expense_types
from tortobaza.finance_db import current_finance_db

SHEET_NAME = "Операции"
BANK_NAME = "TBC"
COLUMN_GROUPS = (
    ("Дата",),
    ("Описание операции", "Назначение"),
    ("Дополнительная информация",),
    ("Списание, GEL", "Расход, GEL"),
    ("Поступление, GEL", "Приход, GEL"),
    ("Остаток, GEL",),
)
DATE_COLUMN = COLUMN_GROUPS[0][0]
OPERATION_COLUMN = COLUMN_GROUPS[1][0]
ADDITIONAL_COLUMN = COLUMN_GROUPS[2][0]
DEBIT_COLUMN = COLUMN_GROUPS[3][0]
CREDIT_COLUMN = COLUMN_GROUPS[4][0]


@dataclass
class ParsedTbcRow:
    date: date
    amount: Decimal
    kind: str
    income_type: str
    counterparty_name: str
    counterparty_iban: str
    description: str
    external_id: str


@dataclass
class TbcStatementImportResult:
    created: int
    skipped: int
    error: str


def import_tbc_statement(file) -> TbcStatementImportResult:
    rows, error = parse_tbc_statement(file)
    if error:
        return TbcStatementImportResult(created=0, skipped=0, error=error)
    if not rows:
        return TbcStatementImportResult(created=0, skipped=0, error="")

    accounts = list(FinancialAccount.objects.filter(bank_name=BANK_NAME))
    if len(accounts) == 0:
        return TbcStatementImportResult(created=0, skipped=0, error=_("TBC account not found"))
    if len(accounts) > 1:
        return TbcStatementImportResult(
            created=0,
            skipped=0,
            error=_("TBC account matched more than once"),
        )
    account = accounts[0]

    external_ids = [row.external_id for row in rows]
    with transaction.atomic(using=current_finance_db()):
        existing = {
            item.external_id: item
            for item in FinancialTransaction.objects.filter(
                account=account,
                external_id__in=external_ids,
            )
        }
        seen = set(existing)
        to_create: list[FinancialTransaction] = []
        to_update: list[FinancialTransaction] = []
        skipped = 0
        for row in rows:
            if row.external_id in seen:
                current = existing.get(row.external_id)
                if current is not None and current.counterparty_name != row.counterparty_name:
                    current.counterparty_name = row.counterparty_name
                    to_update.append(current)
                skipped += 1
                continue
            seen.add(row.external_id)
            to_create.append(
                FinancialTransaction(
                    account=account,
                    date=row.date,
                    amount=row.amount,
                    kind=row.kind,
                    income_type=row.income_type,
                    counterparty_name=row.counterparty_name,
                    counterparty_iban=row.counterparty_iban,
                    description=row.description,
                    external_id=row.external_id,
                )
            )
        assign_expense_types(to_create)
        FinancialTransaction.objects.bulk_create(to_create)
        if to_update:
            FinancialTransaction.objects.bulk_update(to_update, ["counterparty_name"])
    return TbcStatementImportResult(created=len(to_create), skipped=skipped, error="")


def parse_tbc_statement(file) -> tuple[list[ParsedTbcRow], str]:
    workbook = load_workbook(file, data_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        return [], _("Sheet 'Операции' not found")
    worksheet = workbook[SHEET_NAME]
    header_map: dict[str, int] | None = None
    rows: list[ParsedTbcRow] = []
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        if header_map is None:
            names: dict[str, int] = {}
            for index, value in enumerate(values):
                if isinstance(value, str) and value.strip():
                    names[value.strip()] = index
            resolved = resolve_header(names)
            if resolved is not None:
                header_map, missing_columns = resolved
                if missing_columns:
                    return [], _("Missing columns: %(columns)s") % {"columns": ", ".join(missing_columns)}
            continue
        if not any(is_filled(value) for value in values):
            continue
        row_number = excel_row[0].row
        parsed, error = parse_tbc_row(values, header_map, row_number)
        if error:
            return [], error
        rows.append(parsed)
    if header_map is None:
        return [], _("Statement header not found")
    return rows, ""


def resolve_header(names: dict[str, int]) -> tuple[dict[str, int], list[str]] | None:
    if DATE_COLUMN not in names:
        return None
    credit_names = COLUMN_GROUPS[4]
    if not any(name in names for name in credit_names):
        return None
    header_map: dict[str, int] = {}
    missing: list[str] = []
    for aliases in COLUMN_GROUPS:
        index = next((names[alias] for alias in aliases if alias in names), None)
        if index is None:
            missing.append(aliases[0])
        else:
            header_map[aliases[0]] = index
    return header_map, missing


def parse_tbc_row(
    values: list,
    header_map: dict[str, int],
    row_number: int,
) -> tuple[ParsedTbcRow | None, str]:
    operation = text_value(column_value(values, header_map, OPERATION_COLUMN))
    if operation == "":
        return None, _("Row %(row_number)s: missing operation description") % {"row_number": row_number}
    additional = text_value(column_value(values, header_map, ADDITIONAL_COLUMN))
    if additional == "":
        return None, _("Row %(row_number)s: missing additional information") % {"row_number": row_number}
    parts = [part.strip() for part in additional.split(",")]
    if len(parts) < 3:
        return None, _("Row %(row_number)s: additional information must have three parts") % {
            "row_number": row_number
        }
    counterparty_iban = parts[-1]
    if counterparty_iban == "":
        return None, _("Row %(row_number)s: missing IBAN") % {"row_number": row_number}
    has_debit = is_filled(column_value(values, header_map, DEBIT_COLUMN))
    has_credit = is_filled(column_value(values, header_map, CREDIT_COLUMN))
    if has_debit and has_credit:
        return None, _("Row %(row_number)s: both debit and credit are set") % {"row_number": row_number}
    if not has_debit and not has_credit:
        return None, _("Row %(row_number)s: debit and credit are empty") % {"row_number": row_number}
    if has_credit:
        amount = decimal_amount(column_value(values, header_map, CREDIT_COLUMN))
        if amount <= 0:
            return None, _("Row %(row_number)s: amount sign does not match credit") % {"row_number": row_number}
        kind = FinancialTransaction.KIND_INCOME
        income_type = FinancialTransaction.INCOME_TRANSFER
    else:
        amount = -decimal_amount(column_value(values, header_map, DEBIT_COLUMN))
        if amount >= 0:
            return None, _("Row %(row_number)s: amount sign does not match debit") % {"row_number": row_number}
        kind = FinancialTransaction.KIND_EXPENSE
        income_type = ""
    parsed_date = parse_date(column_value(values, header_map, DATE_COLUMN))
    if parsed_date is None:
        return None, _("Row %(row_number)s: invalid date") % {"row_number": row_number}
    description = f"{operation} | {additional}"
    external_id = f"{parsed_date.isoformat()}|{operation}|{additional}|{amount}"
    max_length = FinancialTransaction._meta.get_field("external_id").max_length
    if len(external_id) > max_length:
        return None, _("Row %(row_number)s: external id is too long") % {"row_number": row_number}
    counterparty_name = parts[0]
    if kind == FinancialTransaction.KIND_EXPENSE:
        merchant_name, separator, details = operation.partition(",")
        if separator and "GEL" in details.split(",", 1)[0]:
            counterparty_name = merchant_name.strip()
    return (
        ParsedTbcRow(
            date=parsed_date,
            amount=amount,
            kind=kind,
            income_type=income_type,
            counterparty_name=counterparty_name,
            counterparty_iban=counterparty_iban,
            description=description,
            external_id=external_id,
        ),
        "",
    )
