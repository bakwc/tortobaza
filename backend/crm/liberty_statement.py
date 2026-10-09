from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from crm.bog_statement import decimal_amount, is_filled, parse_date, text_value
from crm.models import FinancialAccount, FinancialTransaction
from crm.transaction_rules import assign_expense_types

SUMMARY_SHEET = "Summary"
STATEMENT_SHEET = "account_statement"
ACCOUNT_LABEL = "Account No:"
REQUIRED_COLUMNS = (
    "Date",
    "Description",
    "Paid Out",
    "Paid In",
    "Document Number",
    "Partner's Name",
    "Partner's Account",
)
ONLINE_PREFIX = "Reimbursement/"
TERMINAL_PREFIX = "ბარათით განაღდებები"


@dataclass
class ParsedLibertyRow:
    account_iban: str
    date: date
    amount: Decimal
    kind: str
    income_type: str
    counterparty_name: str
    counterparty_iban: str
    description: str
    external_id: str


@dataclass
class LibertyStatementImportResult:
    created: int
    skipped: int
    error: str


def import_liberty_statement(file) -> LibertyStatementImportResult:
    rows, error = parse_liberty_statement(file)
    if error:
        return LibertyStatementImportResult(created=0, skipped=0, error=error)

    ibans = {row.account_iban for row in rows}
    found = list(FinancialAccount.objects.filter(iban__in=ibans))
    accounts: dict[str, FinancialAccount] = {}
    ambiguous: list[str] = []
    for account in found:
        if account.iban in accounts:
            ambiguous.append(account.iban)
        else:
            accounts[account.iban] = account
    missing = sorted(ibans - set(accounts))
    ambiguous = sorted(set(ambiguous))
    messages: list[str] = []
    if missing:
        messages.append(_("Unknown accounts: %(accounts)s") % {"accounts": ", ".join(missing)})
    if ambiguous:
        messages.append(_("Accounts matched more than once: %(accounts)s") % {"accounts": ", ".join(ambiguous)})
    if messages:
        return LibertyStatementImportResult(created=0, skipped=0, error="; ".join(messages))

    external_ids = [row.external_id for row in rows]
    account_ids = [account.id for account in accounts.values()]
    with transaction.atomic():
        existing = set(
            FinancialTransaction.objects.filter(
                account_id__in=account_ids,
                external_id__in=external_ids,
            ).values_list("account_id", "external_id")
        )
        to_create: list[FinancialTransaction] = []
        skipped = 0
        for row in rows:
            account = accounts[row.account_iban]
            key = (account.id, row.external_id)
            if key in existing:
                skipped += 1
                continue
            existing.add(key)
            to_create.append(
                FinancialTransaction(
                    account=account,
                    date=row.date,
                    amount=row.amount,
                    kind=row.kind,
                    income_type=row.income_type,
                    counterparty_name=row.counterparty_name,
                    counterparty_iban=row.counterparty_iban,
                    description=row.description,
                    external_id=row.external_id,
                )
            )
        assign_expense_types(to_create)
        FinancialTransaction.objects.bulk_create(to_create)
    return LibertyStatementImportResult(created=len(to_create), skipped=skipped, error="")


def parse_liberty_statement(file) -> tuple[list[ParsedLibertyRow], str]:
    workbook = load_workbook(file, data_only=False)
    if SUMMARY_SHEET not in workbook.sheetnames:
        return [], _("Sheet 'Summary' not found")
    if STATEMENT_SHEET not in workbook.sheetnames:
        return [], _("Sheet 'account_statement' not found")
    account_iban, error = read_account_number(workbook[SUMMARY_SHEET])
    if error:
        return [], error
    worksheet = workbook[STATEMENT_SHEET]
    header_map: dict[str, int] | None = None
    rows: list[ParsedLibertyRow] = []
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        if header_map is None:
            names: dict[str, int] = {}
            for index, value in enumerate(values):
                if isinstance(value, str) and value.strip():
                    names[value.strip()] = index
            if "Date" in names and "Paid In" in names:
                missing_columns = [name for name in REQUIRED_COLUMNS if name not in names]
                if missing_columns:
                    return [], _("Missing columns: %(columns)s") % {"columns": ", ".join(missing_columns)}
                header_map = names
            continue
        if not any(is_filled(value) for value in values):
            continue
        row_number = excel_row[0].row
        parsed, error = parse_liberty_row(values, header_map, row_number, account_iban)
        if error:
            return [], error
        rows.append(parsed)
    if header_map is None:
        return [], _("Statement header not found")
    return rows, ""


def parse_liberty_row(
    values: list,
    header_map: dict[str, int],
    row_number: int,
    account_iban: str,
) -> tuple[ParsedLibertyRow | None, str]:
    description = text_value(column_value(values, header_map, "Description"))
    if description == "":
        return None, _("Row %(row_number)s: missing Description") % {"row_number": row_number}
    document_number = text_value(column_value(values, header_map, "Document Number"))
    if document_number == "":
        return None, _("Row %(row_number)s: missing Document Number") % {"row_number": row_number}
    external_id = f"{document_number}|{description}"
    max_length = FinancialTransaction._meta.get_field("external_id").max_length
    if len(external_id) > max_length:
        return None, _("Row %(row_number)s: external id is too long") % {"row_number": row_number}
    has_paid_out = is_filled(column_value(values, header_map, "Paid Out"))
    has_paid_in = is_filled(column_value(values, header_map, "Paid In"))
    if has_paid_out and has_paid_in:
        return None, _("Row %(row_number)s: both paid out and paid in are set") % {"row_number": row_number}
    if not has_paid_out and not has_paid_in:
        return None, _("Row %(row_number)s: paid out and paid in are empty") % {"row_number": row_number}
    if has_paid_in:
        amount = decimal_amount(column_value(values, header_map, "Paid In"))
        if amount <= 0:
            return None, _("Row %(row_number)s: amount sign does not match paid in") % {"row_number": row_number}
        kind = FinancialTransaction.KIND_INCOME
    else:
        amount = -decimal_amount(column_value(values, header_map, "Paid Out"))
        if amount >= 0:
            return None, _("Row %(row_number)s: amount sign does not match paid out") % {"row_number": row_number}
        kind = FinancialTransaction.KIND_EXPENSE
    parsed_date = parse_date(column_value(values, header_map, "Date"))
    if parsed_date is None:
        return None, _("Row %(row_number)s: invalid date") % {"row_number": row_number}
    if kind == FinancialTransaction.KIND_INCOME and description.startswith(ONLINE_PREFIX):
        income_type = FinancialTransaction.INCOME_ONLINE
    elif kind == FinancialTransaction.KIND_INCOME and description.startswith(TERMINAL_PREFIX):
        income_type = FinancialTransaction.INCOME_TERMINAL
    else:
        income_type = ""
    return (
        ParsedLibertyRow(
            account_iban=account_iban,
            date=parsed_date,
            amount=amount,
            kind=kind,
            income_type=income_type,
            counterparty_name=text_value(column_value(values, header_map, "Partner's Name")),
            counterparty_iban=text_value(column_value(values, header_map, "Partner's Account")),
            description=description,
            external_id=external_id,
        ),
        "",
    )


def read_account_number(worksheet) -> tuple[str, str]:
    found: list[str] = []
    for excel_row in worksheet.iter_rows():
        values = [cell.value for cell in excel_row]
        for index, value in enumerate(values):
            if isinstance(value, str) and value.strip() == ACCOUNT_LABEL:
                if index + 1 >= len(values):
                    found.append("")
                else:
                    found.append(text_value(values[index + 1]))
    numbers = sorted({number for number in found if number})
    if len(numbers) == 1 and "" not in found:
        return numbers[0], ""
    if len(numbers) > 1:
        return "", _("Multiple account numbers: %(numbers)s") % {"numbers": ", ".join(numbers)}
    return "", _("Account No not found")


def column_value(values: list, header_map: dict[str, int], name: str):
    index = header_map[name]
    if index >= len(values):
        return None
    return values[index]
